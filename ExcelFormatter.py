"""フォルダ内（サブフォルダを含む）の .xlsx を一括整形します。

実行例:
    python ExcelFormatter.py "C:\\Users\\User\\Documents\\Excel"

* 各ワークシートの B1 に拡張子を除いたファイル名が含まれる場合:
  横位置=左、縦位置=中央、折り返す、1～2行目を固定
* 含まれない場合:
  同じ配置設定を行い、1行目のみを固定
* 元のファイルは上書きされます。初回実行時のバックアップは
  対象フォルダの「_xlsx_backup」以下に相対パスを維持して保存します。

注意: openpyxl は図形等の Excel の一部要素を保持できない場合があります。
"""

import os
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

BACKUP_DIR_NAME = "_xlsx_backup"
CELL_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def collect_files(root: Path) -> tuple[list[Path], list[Path]]:
    """対象フォルダを再帰的に走査し、フォルダと .xlsx の一覧を返す。"""
    folders: list[Path] = []
    files: list[Path] = []

    for current, directory_names, filenames in os.walk(root):
        # 生成したバックアップの再処理を防止
        directory_names[:] = sorted(
            name for name in directory_names if name != BACKUP_DIR_NAME
        )
        current_path = Path(current)
        folders.extend(current_path / name for name in directory_names)

        for filename in sorted(filenames):
            if filename.startswith("~$"):
                continue  # Excel が編集中に作る一時ファイル
            if Path(filename).suffix.lower() == ".xlsx":
                files.append(current_path / filename)

    return folders, files


def format_file(file: Path, root: Path) -> list[tuple[str, bool, str]]:
    """1つの xlsx を整形して保存。結果は (シート名, B1一致, 固定位置)。"""
    workbook = load_workbook(
        file,
        keep_links=True,  # 外部ブックへのリンクを保持する
        rich_text=True,   # リッチテキストを保持する（openpyxl 3.1.5 対応）
    )

    temp_path: Path | None = None
    results: list[tuple[str, bool, str]] = []

    try:
        for sheet in workbook.worksheets:
            b1 = sheet["B1"].value
            filename_is_in_b1 = isinstance(b1, str) and file.stem in b1

            # 使用範囲内の全セル（空欄のセルを含む）の配置を変更
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = CELL_ALIGNMENT

            # A3: 上2行固定、A2: 上1行固定
            sheet.freeze_panes = "A3" if filename_is_in_b1 else "A2"
            results.append((sheet.title, filename_is_in_b1, sheet.freeze_panes))

        # まず同じディレクトリの一時ファイルへ保存し、保存失敗時の元ファイル破損を防ぐ
        handle, temp_name = tempfile.mkstemp(
            prefix="._xlsx_formatter_", suffix=".xlsx", dir=file.parent
        )
        os.close(handle)
        temp_path = Path(temp_name)
        workbook.save(temp_path)
        workbook.close()

        # 初回のみ元ファイルをバックアップ。再実行時も初回のバックアップは維持
        backup_file = root / BACKUP_DIR_NAME / file.relative_to(root)
        backup_file.parent.mkdir(parents=True, exist_ok=True)
        if not backup_file.exists():
            shutil.copy2(file, backup_file)

        # 一時ファイルを元ファイルに置換
        os.replace(temp_path, file)
        temp_path = None
        return results

    finally:
        workbook.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def main() -> int:
    if len(sys.argv) >= 2:
        folder_input = sys.argv[1]
    else:
        folder_input = input("処理対象フォルダのパスを入力してください: ")

    root = Path(folder_input.strip().strip('"')).expanduser().resolve()
    if not root.is_dir():
        print(f"[ERROR] 指定フォルダが存在しません: {root}")
        return 1

    folders, files = collect_files(root)
    print(f"\n対象フォルダ: {root}")
    print(f"検出したサブフォルダ: {len(folders)}件")
    for folder in folders:
        print(f"  [DIR]  {folder.relative_to(root)}")

    print(f"対象の .xlsx ファイル: {len(files)}件\n")
    if not files:
        print("対象ファイルはありませんでした。")
        return 0

    success = 0
    failure = 0
    for file in files:
        relative = file.relative_to(root)
        try:
            sheet_results = format_file(file, root)
            print(f"[OK] {relative}")
            for title, matched, fixed in sheet_results:
                print(
                    f"     シート「{title}」: B1={'一致' if matched else '不一致'}"
                    f" → {('上2行' if matched else '上1行')}固定 ({fixed})"
                )
            success += 1
        except Exception as exc:
            print(f"[ERROR] {relative}: {exc}")
            failure += 1

    print(f"\n処理完了: 成功 {success}件 / 失敗 {failure}件")
    print(f"元ファイルの初回バックアップ: {root / BACKUP_DIR_NAME}")
    return 0 if failure == 0 else 1


if __name__ == "__main__":
    sys.exit(main())